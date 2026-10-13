"""Phase 6 content: modules 6.1-6.5 (general automation)."""

from auto_support import FAKE_API, needs

API_FILES = {"fake_controller.py": FAKE_API}
REQ = needs("requests") + " and starts a local HTTP server thread"
HTTPX = needs("httpx") + " and starts a local HTTP server thread"
FASTAPI = needs("fastapi httpx")
PW = needs("playwright") + " plus `playwright install chromium`"
PANDAS = needs("pandas openpyxl")
MPL = needs("matplotlib")
PLOTLY = needs("plotly")

ADMIN_HTML = """<!doctype html>
<html><head><title>AP-OS Admin</title>
<style>body{font-family:sans-serif;margin:24px;width:520px} label{display:block;margin:8px 0}
#settings{display:none} #msg{font-weight:bold} .err{color:#b00}</style></head>
<body>
<h1>AP-OS 7.4.1 &middot; ap-lab-07</h1>
<form id="login">
  <label>Username <input id="user"></label>
  <label>Password <input id="pass" type="password"></label>
  <button type="submit">Sign in</button>
</form>
<section id="settings">
  <p id="status">Loading status...</p>
  <label>Band <select id="band"><option>2G</option><option selected>5G</option><option>6G</option></select></label>
  <label>Channel <input id="channel" value="36"></label>
  <button id="save" type="button">Save</button>
  <p id="msg"></p>
</section>
<script>
document.querySelector("#login").addEventListener("submit", async (ev) => {
  ev.preventDefault();
  document.querySelector("#login").style.display = "none";
  document.querySelector("#settings").style.display = "block";
  const r = await fetch("/api/status");
  const st = document.querySelector("#status");
  if (!r.ok) { st.textContent = "Error: status unavailable (" + r.status + ")"; st.className = "err"; return; }
  const s = await r.json();
  st.textContent = "Clients: " + s.clients + " | Uptime: " + s.uptime;
});
document.querySelector("#save").addEventListener("click", async () => {
  const msg = document.querySelector("#msg");
  msg.textContent = "Saving...";
  const body = {band: document.querySelector("#band").value, channel: Number(document.querySelector("#channel").value)};
  await fetch("/api/radio", {method: "POST", headers: {"Content-Type": "application/json"}, body: JSON.stringify(body)});
  setTimeout(() => { msg.textContent = "Saved: " + body.band + " channel " + body.channel; }, 300);
});
</script>
</body></html>
"""

PW_FILES = {"ap_admin.html": ADMIN_HTML}

M_6_1 = {
    "id": "6.1",
    "title": "REST APIs with requests & httpx",
    "hours": 2.5,
    "summary": "Call device and controller REST APIs: sessions, auth headers, query strings, JSON bodies, status handling, timeouts, retries, pagination and concurrent async requests.",
    "concepts": [
        {
            "title": "requests: sessions, JSON, errors and timeouts",
            "files": API_FILES,
            "local_only": REQ,
            "concept": "`requests` is the de-facto HTTP client. A `Session` keeps connections open and applies shared headers (auth, Accept) to every call. `params=` builds the query string, `json=` serialises the body and sets `Content-Type`, `r.json()` parses the reply. Non-2xx statuses are **not** exceptions until you call `r.raise_for_status()`, and there is **no default timeout**, so always pass `timeout=`.",
            "why": "Wi-Fi controllers, cloud management portals, traffic generators and ticketing systems all expose REST APIs. Most of a test engineer's integration glue is exactly these few calls, done safely.",
            "syntax": "with requests.Session() as s:\n    s.headers[\"Authorization\"] = f\"Bearer {token}\"\n    r = s.get(url, params={...}, timeout=5)\n    r.raise_for_status()\n    data = r.json()",
            "code": """import requests  ## Third-party HTTP client: pip install requests

from fake_controller import TOKEN, start  ## Simulated controller API (see supporting files)

BASE = start()  ## Base URL like http://127.0.0.1:54321

with requests.Session() as s:  ## Reuses TCP connections and shared settings
    s.headers.update({"Authorization": f"Bearer {TOKEN}", "Accept": "application/json"})  ## Sent with every request
    r = s.get(f"{BASE}/api/v1/aps", params={"band": "5G"}, timeout=5)  ## ?band=5G is built for you
    r.raise_for_status()  ## Raises HTTPError for 4xx/5xx, does nothing for 2xx
    print(r.status_code, r.headers["Content-Type"])  ## Status and a response header
    print([ap["name"] for ap in r.json()])  ## Parsed JSON body (a list here)

    r = s.post(f"{BASE}/api/v1/aps/ap-3/channel", json={"band": "5G", "channel": 165}, timeout=5)  ## json= serialises and sets Content-Type
    print(r.status_code, r.json())  ## The updated resource

    r = s.post(f"{BASE}/api/v1/aps/ap-3/channel", json={"band": "5G", "channel": 52}, timeout=5)  ## Channel the API rejects
    try:  ## Turn the error status into an exception
        r.raise_for_status()  ## 422 -> requests.HTTPError
    except requests.HTTPError as e:  ## The response is attached to the exception
        print("HTTP", e.response.status_code, e.response.json()["error"])  ## Server's explanation

    try:  ## Timeouts are exceptions too
        s.get(f"{BASE}/api/v1/slow", timeout=0.5)  ## This endpoint takes 2 s
    except requests.Timeout:  ## Covers both connect and read timeouts
        print("timed out after 0.5 s")  ## Fail fast instead of hanging a CI job""",
            "under": "`requests` builds on `urllib3`, which pools connections per host inside the Session, so repeated calls skip the TCP (and TLS) handshake. `timeout=5` is applied separately to connecting and to each socket read, not to the whole download. `r.json()` raises `requests.JSONDecodeError` if the body isn't JSON, which is common with HTML error pages from proxies.",
            "gotchas": [
                {
                    "title": "Using the body without checking the status",
                    "files": API_FILES,
                    "local_only": REQ,
                    "note": "A 404 still has a JSON body (the error message), so code that skips the status check carries on with nonsense: here `channel` is silently None. Check the status first.",
                    "wrong": """import requests
from fake_controller import start

BASE = start()
ap = requests.get(f"{BASE}/api/v1/aps/ap-7", timeout=5).json()
print("channel:", ap.get("channel"))""",
                    "fixed": """import requests
from fake_controller import start

BASE = start()
r = requests.get(f"{BASE}/api/v1/aps/ap-7", timeout=5)
try:
    r.raise_for_status()
    print("channel:", r.json()["channel"])
except requests.HTTPError:
    print("lookup failed:", r.status_code, r.json()["error"])""",
                },
                {
                    "title": "data= instead of json=",
                    "files": API_FILES,
                    "local_only": REQ,
                    "note": "`data={...}` sends a form-encoded body (`band=5G&channel=36`), not JSON, so JSON APIs reject it (415 or 422). Use `json=`.",
                    "wrong": """import requests
from fake_controller import TOKEN, start

BASE = start()
r = requests.post(f"{BASE}/api/v1/aps/ap-1/channel", data={"band": "5G", "channel": 36},
                  headers={"Authorization": f"Bearer {TOKEN}"}, timeout=5)
print(r.status_code, r.json())""",
                    "fixed": """import requests
from fake_controller import TOKEN, start

BASE = start()
r = requests.post(f"{BASE}/api/v1/aps/ap-1/channel", json={"band": "5G", "channel": 36},
                  headers={"Authorization": f"Bearer {TOKEN}"}, timeout=5)
print(r.status_code, r.json())""",
                },
            ],
            "best": "Wrap each API in a small client class around one Session, set `timeout` on every call, call `raise_for_status()` (or check `r.ok`) before using the body, keep tokens in environment variables, and log method, URL and status (never the Authorization header).",
        },
        {
            "title": "Retries, pagination and httpx (sync and async)",
            "files": API_FILES,
            "local_only": HTTPX,
            "concept": "`httpx` has a requests-like API plus `base_url`, a default 5-second timeout, HTTP/2 and an `AsyncClient`. Retry only transient failures (429, 502, 503, 504, connection errors), never a 4xx caused by your request. APIs that return large collections page them: follow the `next` link (or offset/cursor) until there isn't one. For many independent calls, an `AsyncClient` with `asyncio.gather` sends them concurrently.",
            "why": "Controllers get busy during firmware rollouts and answer 503, inventory APIs return clients 100 at a time, and polling 200 APs one by one takes minutes while concurrent requests take seconds.",
            "syntax": "with httpx.Client(base_url=BASE, timeout=5.0) as c:\n    c.get(\"/path\").raise_for_status().json()\n\nasync with httpx.AsyncClient(base_url=BASE) as c:\n    rs = await asyncio.gather(*(c.get(p) for p in paths))",
            "code": """import asyncio  ## For the async client

import httpx  ## pip install httpx

from fake_controller import start  ## Simulated controller API

BASE = start()  ## Fresh server for this example
TRANSIENT = {429, 502, 503, 504}  ## Statuses worth retrying


def get_with_retry(client, path, attempts=4):  ## Retry only what can succeed later
    for attempt in range(1, attempts + 1):  ## 1-based for the log line
        r = client.get(path)  ## base_url is prepended automatically
        if r.status_code not in TRANSIENT:  ## Success, or a permanent error
            return r  ## Stop retrying
        print(f"attempt {attempt}: {r.status_code}, retrying")  ## Real code would sleep with backoff and jitter here
    return r  ## Last response after giving up


def all_clients(client):  ## Follow "next" links until the last page
    path, items = "/api/v1/clients", []  ## Start at page 1
    while path:  ## next is None on the last page
        page = client.get(path).raise_for_status().json()  ## In httpx, raise_for_status() returns the response
        items += page["items"]  ## Collect this page's items
        path = page["next"]  ## Relative link to the following page
    return items  ## Everything, across all pages


with httpx.Client(base_url=BASE, timeout=5.0) as client:  ## One connection pool for all calls
    print(get_with_retry(client, "/api/v1/flaky").json())  ## Busy twice, then succeeds
    print(len(all_clients(client)), "clients across pages")  ## 3 per page on this API


async def channels(names):  ## Many independent requests at once
    async with httpx.AsyncClient(base_url=BASE, timeout=5.0) as client:  ## Async version of the client
        replies = await asyncio.gather(*(client.get(f"/api/v1/aps/{n}") for n in names))  ## All requests in flight together
        return {n: r.json().get("channel", r.status_code) for n, r in zip(names, replies)}  ## Channel, or the status for errors


print(asyncio.run(channels(["ap-1", "ap-2", "ap-9"])))  ## ap-9 doesn't exist, so it maps to 404""",
            "under": "`asyncio.gather` starts every coroutine and waits for all of them; the requests overlap while each waits for the network. Results come back in the order you passed them, not the order they finished. For real retries add exponential backoff with jitter (Phase 4.5) or use `tenacity`; for requests, `urllib3.util.Retry` mounted through an `HTTPAdapter` does it declaratively.",
            "gotchas": [],
            "best": "Retry idempotent requests (GET, PUT) on transient errors with backoff and a cap; be careful retrying POSTs that create things. Limit concurrency (an `asyncio.Semaphore` or httpx `Limits`) so you don't overload the controller you're testing, and respect `Retry-After` headers on 429.",
        },
    ],
    "exercises": [
        {
            "level": "Easy",
            "files": API_FILES,
            "local_only": REQ,
            "prompt": "List every AP from `/api/v1/aps` and print `name: clients`, busiest first.",
            "solution": "One GET with a timeout, a status check, then sort the parsed list by the clients field in reverse.",
            "solution_code": """import requests
from fake_controller import start

BASE = start()
r = requests.get(f"{BASE}/api/v1/aps", timeout=5)
r.raise_for_status()
for ap in sorted(r.json(), key=lambda a: a["clients"], reverse=True):
    print(f"{ap['name']}: {ap['clients']}")""",
        },
        {
            "level": "Medium",
            "files": API_FILES,
            "local_only": REQ,
            "prompt": "Make a requests Session retry 503 responses automatically (up to 3 retries, no delay for the demo) using urllib3's `Retry` and an `HTTPAdapter`, then call `/api/v1/flaky` once.",
            "solution": "Mount an HTTPAdapter configured with Retry(status_forcelist=[503]) on the http:// prefix; the session hides the two failures, and the body shows it succeeded on attempt 3.",
            "solution_code": """import requests
from requests.adapters import HTTPAdapter
from urllib3.util import Retry
from fake_controller import start

BASE = start()
retry = Retry(total=3, status_forcelist=[503], backoff_factor=0, allowed_methods=["GET"])
with requests.Session() as s:
    s.mount("http://", HTTPAdapter(max_retries=retry))
    r = s.get(f"{BASE}/api/v1/flaky", timeout=5)
    print(r.status_code, r.json())""",
        },
        {
            "level": "Hard",
            "files": API_FILES,
            "local_only": REQ,
            "prompt": "Write a `ControllerClient` class with `list_aps(band=None)` and `set_channel(name, band, channel)`. API errors must raise your own `ControllerError` carrying the status code and the server's message.",
            "solution": "Encapsulate the Session, base URL, token and timeout; convert HTTP errors into a domain exception so callers never see raw responses.",
            "solution_code": """import requests
from fake_controller import TOKEN, start


class ControllerError(Exception):
    def __init__(self, status, message):
        super().__init__(f"{status}: {message}")
        self.status = status


class ControllerClient:
    def __init__(self, base, token, timeout=5):
        self.base, self.timeout = base.rstrip("/"), timeout
        self.s = requests.Session()
        self.s.headers["Authorization"] = f"Bearer {token}"

    def _call(self, method, path, **kw):
        r = self.s.request(method, f"{self.base}/api/v1{path}", timeout=self.timeout, **kw)
        if not r.ok:
            raise ControllerError(r.status_code, r.json().get("error", r.text))
        return r.json()

    def list_aps(self, band=None):
        return self._call("GET", "/aps", params={"band": band} if band else None)

    def set_channel(self, name, band, channel):
        return self._call("POST", f"/aps/{name}/channel", json={"band": band, "channel": channel})


api = ControllerClient(start(), TOKEN)
print([ap["name"] for ap in api.list_aps("6G")])
print(api.set_channel("ap-2", "6G", 37)["channel"])
try:
    api.set_channel("ap-2", "6G", 38)
except ControllerError as e:
    print("ControllerError", e.status, e)""",
        },
    ],
    "interview": [
        {"level": "Beginner", "q": "What's the difference between `params`, `data` and `json` in requests?", "a": "params builds the URL query string. data sends a form-encoded (or raw) body. json serialises a Python object to JSON and sets Content-Type: application/json.", "follow": "What happens if you pass both data and json? (data wins; json is ignored.)"},
        {"level": "Intermediate", "q": "Why should every request have a timeout?", "a": "requests has no default timeout, so a dead peer can hang the process forever, stalling CI jobs. Set connect and read timeouts (e.g. timeout=(3, 30)) and handle requests.Timeout.", "follow": "Is the read timeout a limit on total download time? (No, it's per socket read.)"},
        {"level": "Intermediate", "q": "Which failures would you retry, and how?", "a": "Transient ones: connection errors, 429, 502, 503, 504, preferably only for idempotent methods. Use exponential backoff with jitter, a max attempt count, and honour Retry-After. Never retry 400/401/404/422, they won't change.", "follow": "Why is retrying a POST risky? (It might create a duplicate resource.)"},
        {"level": "Advanced", "q": "requests vs httpx vs aiohttp?", "a": "requests: mature, sync only. httpx: requests-like API, sync and async, HTTP/2, default timeouts, good for both styles. aiohttp: async only, very fast, also a server framework. For a mostly sync automation codebase with occasional fan-out, httpx is a good default.", "follow": "How do you limit concurrency with httpx async? (httpx.Limits or an asyncio.Semaphore.)"},
    ],
    "quiz": [
        {"q": "`requests.get(url)` returns 404. What happens?", "options": ["Raises HTTPError", "Returns a Response with status 404", "Returns None", "Retries"], "answer": 1, "why": "Only raise_for_status() raises."},
        {"q": "Default timeout in requests:", "options": ["5 s", "30 s", "None (wait forever)", "60 s"], "answer": 2, "why": "Always pass timeout=."},
        {"q": "Which status is safe to retry?", "options": ["400", "401", "503", "422"], "answer": 2, "why": "Service unavailable is transient."},
        {"q": "Benefit of a Session:", "options": ["Async", "Connection reuse and shared headers", "Automatic retries", "JSON parsing"], "answer": 1, "why": "Pooling and shared settings."},
        {"q": "asyncio.gather returns results in...", "options": ["Completion order", "The order passed in", "Random order", "Sorted order"], "answer": 1, "why": "Input order."},
    ],
    "cards": [
        {"front": "Safe GET", "back": "`r = s.get(url, params=p, timeout=5); r.raise_for_status(); r.json()`"},
        {"front": "JSON body", "back": "`json=payload` (not data=)"},
        {"front": "Retryable statuses", "back": "429, 502, 503, 504 (+ connection errors), idempotent methods"},
        {"front": "Pagination", "back": "Loop while page['next']: fetch, extend, follow"},
        {"front": "Concurrent requests", "back": "`httpx.AsyncClient` + `asyncio.gather(...)`"},
    ],
}


BENCH_API = '''from fastapi import Depends, FastAPI, Header, HTTPException

app = FastAPI(title="Bench API")


def get_inventory() -> dict:
    raise RuntimeError("would query the real CMDB")


def require_token(authorization: str = Header(default="")):
    if authorization != "Bearer lab-token":
        raise HTTPException(status_code=401, detail="bad or missing token")


@app.get("/aps/{name}", dependencies=[Depends(require_token)])
def get_ap(name: str, inventory: dict = Depends(get_inventory)):
    if name not in inventory:
        raise HTTPException(status_code=404, detail=f"no AP {name}")
    return inventory[name]
'''

M_6_2 = {
    "id": "6.2",
    "title": "Building APIs with FastAPI",
    "hours": 2.5,
    "summary": "Build small HTTP services for your lab (bench status, test triggers, results) with FastAPI: typed routes, pydantic validation, errors, dependencies, and testing with TestClient.",
    "concepts": [
        {
            "title": "Routes, validation and TestClient",
            "local_only": FASTAPI,
            "concept": "FastAPI turns type-hinted functions into HTTP endpoints. Path parameters come from `{name}` in the route, query parameters from other plain arguments, and the body from a pydantic model, all validated automatically: bad input gets a 422 with one error per field. Raise `HTTPException` for errors. `TestClient` calls the app in-process, so you can test it without starting a server.",
            "why": "Test teams build small services all the time: a bench reservation API, a results collector, a webhook receiver for CI. FastAPI gives validation, docs (`/docs`) and a client-friendly JSON API with very little code.",
            "syntax": "app = FastAPI()\n\n@app.get(\"/aps/{name}\")\ndef get_ap(name: str, verbose: bool = False): ...\n\n$ uvicorn bench_api:app --reload",
            "code": """from fastapi import FastAPI, HTTPException, Query  ## pip install fastapi
from fastapi.testclient import TestClient  ## In-process client (uses httpx)
from pydantic import BaseModel, Field  ## Request body models

app = FastAPI(title="Bench API")  ## The application object that uvicorn serves
APS = {"ap-1": {"name": "ap-1", "band": "5G", "channel": 36}}  ## In-memory store for the example


class ChannelChange(BaseModel):  ## Body schema, validated before your code runs
    band: str = Field(pattern="^(2G|5G|6G)$")  ## Regex constraint
    channel: int = Field(ge=1, le=233)  ## Range constraint


@app.get("/aps")  ## GET /aps
def list_aps(band: str | None = Query(default=None)):  ## Optional ?band= query parameter
    return [ap for ap in APS.values() if band in (None, ap["band"])]  ## Returned objects are sent as JSON


@app.get("/aps/{name}")  ## {name} is a path parameter
def get_ap(name: str):  ## Filled from the URL
    if name not in APS:  ## Unknown AP
        raise HTTPException(status_code=404, detail=f"no AP {name}")  ## Becomes a 404 JSON response
    return APS[name]  ## 200 with the AP


@app.put("/aps/{name}/channel")  ## Update endpoint
def set_channel(name: str, change: ChannelChange):  ## Body parsed into the model
    ap = get_ap(name)  ## Reuse the 404 logic
    ap.update(change.model_dump())  ## Apply the validated fields
    return ap  ## Updated AP


client = TestClient(app)  ## Calls the app directly: no server, no port
print(client.get("/aps", params={"band": "5G"}).json())  ## Filtered list
print(client.put("/aps/ap-1/channel", json={"band": "6G", "channel": 37}).json())  ## Valid update
r = client.put("/aps/ap-1/channel", json={"band": "7G", "channel": 999})  ## Two invalid fields
print(r.status_code, [e["loc"][-1] for e in r.json()["detail"]])  ## 422, one error per bad field
print(client.get("/aps/ap-9").status_code, client.get("/aps/ap-9").json())  ## The HTTPException response""",
            "under": "At startup FastAPI inspects each function's signature: simple types become path or query parameters, pydantic models become the JSON body. Requests are validated against generated pydantic models before your function runs, and the same information produces the OpenAPI schema behind `/docs`. Plain `def` endpoints run in a thread pool; `async def` endpoints run on the event loop and must not block.",
            "gotchas": [
                {
                    "title": "Route order: a path parameter swallows a fixed path",
                    "local_only": FASTAPI,
                    "note": "Routes are matched in the order they're declared. `/aps/{name}` declared first also matches `/aps/summary`, with name=\"summary\". Declare fixed paths before parameterised ones.",
                    "wrong": """from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

app = FastAPI()
APS = {"ap-1": {"band": "5G"}}


@app.get("/aps/{name}")
def get_ap(name: str):
    if name not in APS:
        raise HTTPException(404, f"no AP {name}")
    return APS[name]


@app.get("/aps/summary")
def summary():
    return {"count": len(APS)}


print(TestClient(app).get("/aps/summary").json())""",
                    "fixed": """from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient

app = FastAPI()
APS = {"ap-1": {"band": "5G"}}


@app.get("/aps/summary")
def summary():
    return {"count": len(APS)}


@app.get("/aps/{name}")
def get_ap(name: str):
    if name not in APS:
        raise HTTPException(404, f"no AP {name}")
    return APS[name]


print(TestClient(app).get("/aps/summary").json())""",
                },
            ],
            "best": "Model every request and response with pydantic, return proper status codes (201 for created, 202 for accepted jobs, 404, 409 for conflicts), keep endpoint functions thin (call into a service layer), and test every route with TestClient.",
        },
        {
            "title": "Dependencies, auth and testing with overrides",
            "runner": "pytest",
            "filename": "test_bench_api.py",
            "files": {"bench_api.py": BENCH_API, "pytest.ini": "[pytest]\nfilterwarnings =\n    ignore::DeprecationWarning:starlette.*\n"},
            "args": ["-v"],
            "concept": "`Depends(fn)` injects the result of `fn` into an endpoint: a database session, the current user, a token check. Dependencies can themselves read headers (`Header()`) and raise HTTPException. In tests, `app.dependency_overrides[real] = fake` swaps a dependency for the whole app, so you test the HTTP layer without the real database or CMDB.",
            "why": "The bench API needs auth on every write and needs data from an inventory system you can't reach from CI. Dependencies keep both concerns out of the endpoint code and make them trivial to replace in tests.",
            "syntax": "def require_token(authorization: str = Header(\"\")): ...\n\n@app.get(\"/x\", dependencies=[Depends(require_token)])\ndef x(db = Depends(get_db)): ...\n\napp.dependency_overrides[get_db] = lambda: fake_db",
            "code": """import pytest  ## Test framework
from fastapi.testclient import TestClient  ## In-process HTTP client

from bench_api import app, get_inventory  ## The app and the dependency we'll replace

FAKE_INVENTORY = {"ap-7": {"name": "ap-7", "band": "6G", "channel": 5}}  ## Test data instead of the CMDB


@pytest.fixture  ## Fresh override per test
def client():  ## Yields a client wired to fake data
    app.dependency_overrides[get_inventory] = lambda: FAKE_INVENTORY  ## Every Depends(get_inventory) now gets this
    yield TestClient(app)  ## The test runs here
    app.dependency_overrides.clear()  ## Restore the real dependencies


def test_requires_token(client):  ## No Authorization header
    assert client.get("/aps/ap-7").status_code == 401  ## Rejected by require_token


def test_reads_overridden_inventory(client):  ## Authenticated request
    r = client.get("/aps/ap-7", headers={"Authorization": "Bearer lab-token"})  ## Valid token
    assert r.status_code == 200 and r.json()["channel"] == 5  ## Data came from the override


def test_unknown_ap_is_404(client):  ## Error path
    r = client.get("/aps/ap-1", headers={"Authorization": "Bearer lab-token"})  ## Not in the fake inventory
    assert r.json() == {"detail": "no AP ap-1"}  ## FastAPI's error shape


def test_openapi_documents_route(client):  ## FastAPI generates an OpenAPI schema
    assert "/aps/{name}" in client.get("/openapi.json").json()["paths"]  ## The same data behind /docs""",
            "under": "FastAPI resolves dependencies per request, caching each one within that request, and supports nested dependencies and `yield` dependencies (setup, then cleanup after the response). `dependency_overrides` is a plain dict on the app consulted during resolution, which is why clearing it in teardown matters: overrides would otherwise leak into other tests.",
            "gotchas": [],
            "best": "Put auth, DB sessions and external clients behind dependencies, override them in a fixture (and always clear), and keep a handful of tests that run with the real dependencies against a staging environment.",
        },
    ],
    "exercises": [
        {
            "level": "Easy",
            "local_only": FASTAPI,
            "prompt": "Add a `/health` endpoint returning `{\"status\": \"ok\", \"version\": \"1.4.0\"}` and check it with TestClient.",
            "solution": "A one-line endpoint; TestClient shows status 200 and the JSON.",
            "solution_code": """from fastapi import FastAPI
from fastapi.testclient import TestClient

app = FastAPI()


@app.get("/health")
def health():
    return {"status": "ok", "version": "1.4.0"}


r = TestClient(app).get("/health")
print(r.status_code, r.json())""",
        },
        {
            "level": "Medium",
            "local_only": FASTAPI,
            "prompt": "Create `POST /runs` that accepts `{suite, bench}` (bench must match `lab-\\d+`), returns 201 with a generated id, and `GET /runs` that lists them.",
            "solution": "A pydantic model with a pattern, `status_code=201` on the decorator, and an in-memory list.",
            "solution_code": """from fastapi import FastAPI
from fastapi.testclient import TestClient
from pydantic import BaseModel, Field

app = FastAPI()
RUNS = []


class RunRequest(BaseModel):
    suite: str
    bench: str = Field(pattern=r"^lab-\\d+$")


@app.post("/runs", status_code=201)
def create_run(req: RunRequest):
    run = {"id": len(RUNS) + 1, **req.model_dump(), "state": "queued"}
    RUNS.append(run)
    return run


@app.get("/runs")
def list_runs():
    return RUNS


c = TestClient(app)
print(c.post("/runs", json={"suite": "smoke", "bench": "lab-7"}).status_code)
print(c.post("/runs", json={"suite": "rvr", "bench": "garage"}).status_code)
print(c.get("/runs").json())""",
        },
        {
            "level": "Hard",
            "local_only": FASTAPI,
            "prompt": "Make `POST /runs` return 202 Accepted with a `Location` header, and finish the run in a background task so `GET /runs/{id}` reports `done` afterwards.",
            "solution": "`BackgroundTasks` runs after the response is sent (TestClient runs it before returning). The Response object sets the Location header.",
            "solution_code": """from fastapi import BackgroundTasks, FastAPI, HTTPException, Response
from fastapi.testclient import TestClient

app = FastAPI()
RUNS = {}


def execute(run_id):
    RUNS[run_id]["state"] = "done"
    RUNS[run_id]["passed"] = 41


@app.post("/runs", status_code=202)
def start_run(background: BackgroundTasks, response: Response):
    run_id = len(RUNS) + 1
    RUNS[run_id] = {"id": run_id, "state": "queued"}
    background.add_task(execute, run_id)
    response.headers["Location"] = f"/runs/{run_id}"
    return RUNS[run_id]


@app.get("/runs/{run_id}")
def get_run(run_id: int):
    if run_id not in RUNS:
        raise HTTPException(404, "no such run")
    return RUNS[run_id]


c = TestClient(app)
r = c.post("/runs")
print(r.status_code, r.headers["location"], r.json())
print(c.get(r.headers["location"]).json())""",
        },
    ],
    "interview": [
        {"level": "Beginner", "q": "How does FastAPI know where a parameter comes from?", "a": "Names in the route path are path parameters; other simple-typed arguments are query parameters; pydantic models are the JSON body; Header(), Cookie(), Depends() mark other sources.", "follow": "What status does invalid input return? (422 Unprocessable Entity, with per-field details.)"},
        {"level": "Intermediate", "q": "How do you test a FastAPI app without a server?", "a": "fastapi.testclient.TestClient(app) calls the ASGI app in-process. Use dependency_overrides to swap databases, auth and external clients.", "follow": "How do you test async-only behaviour? (httpx.AsyncClient with ASGITransport.)"},
        {"level": "Intermediate", "q": "def vs async def endpoints?", "a": "def endpoints run in a thread pool, so blocking calls are fine. async def endpoints run on the event loop; a blocking call (requests, time.sleep) inside them stalls every request.", "follow": "How would you call a blocking library from async def? (run_in_threadpool / asyncio.to_thread.)"},
        {"level": "Advanced", "q": "Design an API that triggers long test runs.", "a": "POST /runs validates input and returns 202 with a run id and Location header; a worker queue (Celery, RQ, or a simple DB-backed queue) executes; GET /runs/{id} returns state and links to artifacts; optional webhook or SSE for completion. Idempotency key header to avoid duplicate runs.", "follow": "Why not run the suite inside the request? (Timeouts, worker starvation, no resilience to restarts.)"},
    ],
    "quiz": [
        {"q": "Invalid body fields give status...", "options": ["400", "404", "422", "500"], "answer": 2, "why": "Validation errors are 422."},
        {"q": "`/aps/{name}` declared before `/aps/summary`. GET /aps/summary hits...", "options": ["summary()", "get_ap(name='summary')", "404", "Both"], "answer": 1, "why": "First match wins."},
        {"q": "Swap a dependency in tests:", "options": ["monkeypatch only", "app.dependency_overrides[dep] = fake", "Edit the source", "Depends(fake)"], "answer": 1, "why": "Built-in override mechanism."},
        {"q": "Status for an accepted async job:", "options": ["200", "201", "202", "204"], "answer": 2, "why": "202 Accepted."},
        {"q": "Interactive docs URL:", "options": ["/swagger", "/docs", "/api", "/help"], "answer": 1, "why": "/docs (and /redoc)."},
    ],
    "cards": [
        {"front": "Run a FastAPI app", "back": "`uvicorn module:app --reload`"},
        {"front": "Error response", "back": "`raise HTTPException(status_code=404, detail='...')`"},
        {"front": "Body validation", "back": "pydantic model argument; Field(pattern=..., ge=..., le=...)"},
        {"front": "Inject auth/DB", "back": "`Depends(fn)`; override with `app.dependency_overrides`"},
        {"front": "Test without server", "back": "`TestClient(app).get('/path')`"},
    ],
}


M_6_3 = {
    "id": "6.3",
    "title": "Web UI Automation with Playwright",
    "hours": 2.5,
    "summary": "Drive a device's web admin UI with Playwright: user-facing locators, auto-waiting assertions, network interception to fake the device, Page Objects and screenshots.",
    "concepts": [
        {
            "title": "Locators, actions, expect and route",
            "files": PW_FILES,
            "local_only": PW,
            "concept": "Playwright controls real Chromium, Firefox or WebKit. Find elements with **user-facing locators** (`get_by_label`, `get_by_role(\"button\", name=...)`, `get_by_text`) rather than brittle CSS paths. Actions (`fill`, `click`, `select_option`) wait for the element to be visible and enabled. `expect(locator).to_have_text(...)` retries until it passes or times out. `page.route()` intercepts network requests, so you can serve the UI and fake the device's API from Python.",
            "why": "Many APs, controllers and cloud portals are configured through a web UI, and some settings exist nowhere else. UI tests also catch regressions the API tests can't: a Save button that silently sends the wrong channel.",
            "syntax": "with sync_playwright() as p:\n    page = p.chromium.launch().new_page()\n    page.goto(url)\n    page.get_by_label(\"Channel\").fill(\"149\")\n    page.get_by_role(\"button\", name=\"Save\").click()\n    expect(page.locator(\"#msg\")).to_have_text(\"Saved\")",
            "code": """import json  ## To decode what the page sends
from pathlib import Path  ## To read the page

from playwright.sync_api import expect, sync_playwright  ## pip install playwright && playwright install chromium

HTML = Path("ap_admin.html").read_text()  ## The admin UI under test (supporting file)
sent = []  ## Requests the UI sends to the "device"


def fake_device(route):  ## Handles every request to http://ap.lab/
    url = route.request.url  ## What the browser asked for
    if url.endswith("/api/status"):  ## Status poll after login
        route.fulfill(json={"clients": 12, "uptime": "3d"})  ## Canned JSON reply
    elif url.endswith("/api/radio"):  ## The Save button posts here
        sent.append(json.loads(route.request.post_data))  ## Record the payload
        route.fulfill(json={"ok": True})  ## Pretend the device accepted it
    else:  ## Anything else is the page itself
        route.fulfill(body=HTML, content_type="text/html")  ## Serve the HTML


with sync_playwright() as p:  ## Starts the Playwright driver process
    browser = p.chromium.launch()  ## Headless by default
    page = browser.new_page()  ## New tab with its own cookies and storage
    page.route("http://ap.lab/**", fake_device)  ## Intercept all traffic to the device
    page.goto("http://ap.lab/")  ## Loads the page through our route
    page.get_by_label("Username").fill("admin")  ## Find an input by its visible label
    page.get_by_label("Password").fill("lab123")  ## Same for the password
    page.get_by_role("button", name="Sign in").click()  ## Accessible role and name
    expect(page.get_by_text("Clients: 12")).to_be_visible()  ## Waits until the status has loaded
    page.get_by_label("Band").select_option("6G")  ## Dropdown
    page.get_by_label("Channel").fill("37")  ## Replaces the field's content
    page.get_by_role("button", name="Save").click()  ## Triggers the POST
    expect(page.locator("#msg")).to_have_text("Saved: 6G channel 37")  ## Retries until the text matches
    print(sent)  ## What the UI actually sent to the device
    print(page.title())  ## The document title
    browser.close()  ## Shut the browser down""",
            "under": "Playwright talks to the browser over its DevTools-style protocol. A locator is a recipe, not an element: it's resolved fresh each time it's used, so it survives re-renders. Before each action Playwright runs actionability checks (attached, visible, stable, enabled, receives events); `expect` assertions poll until the condition is true or the timeout (5 s by default) expires. `route.fulfill` answers from Python without any network.",
            "gotchas": [
                {
                    "title": "Reading text right after a click",
                    "files": PW_FILES,
                    "local_only": PW,
                    "note": "The page updates the message asynchronously after the device answers. `text_content()` reads the DOM **now**, so it sees the intermediate state. Use an auto-retrying `expect` (never `time.sleep`).",
                    "wrong": """from pathlib import Path
from playwright.sync_api import sync_playwright

HTML = Path("ap_admin.html").read_text()
with sync_playwright() as p:
    page = p.chromium.launch().new_page()
    page.route("http://ap.lab/**", lambda r: r.fulfill(json={"clients": 1, "uptime": "1h"}) if "/api/" in r.request.url else r.fulfill(body=HTML, content_type="text/html"))
    page.goto("http://ap.lab/")
    page.get_by_role("button", name="Sign in").click()
    page.get_by_role("button", name="Save").click()
    print(repr(page.locator("#msg").text_content()))""",
                    "fixed": """from pathlib import Path
from playwright.sync_api import expect, sync_playwright

HTML = Path("ap_admin.html").read_text()
with sync_playwright() as p:
    page = p.chromium.launch().new_page()
    page.route("http://ap.lab/**", lambda r: r.fulfill(json={"clients": 1, "uptime": "1h"}) if "/api/" in r.request.url else r.fulfill(body=HTML, content_type="text/html"))
    page.goto("http://ap.lab/")
    page.get_by_role("button", name="Sign in").click()
    page.get_by_role("button", name="Save").click()
    expect(page.locator("#msg")).to_have_text("Saved: 5G channel 36")
    print(repr(page.locator("#msg").text_content()))""",
                },
            ],
            "best": "Prefer role and label locators (they double as an accessibility check), assert with `expect`, never sleep, fake slow or unavailable backends with `page.route`, and run headless in CI with tracing on failure (`context.tracing.start(screenshots=True, snapshots=True)`).",
        },
        {
            "title": "Page Objects and screenshots",
            "files": PW_FILES,
            "local_only": PW,
            "capture": "admin.png",
            "image_alt": "Screenshot of the simulated AP admin page after saving channel 149 on 5 GHz",
            "concept": "A **Page Object** wraps one page (or component) in a class: locators are defined once as attributes, and user actions become methods (`login()`, `set_channel()`). Tests then read like the manual test case. `page.screenshot(path=...)` saves a PNG of the page (`full_page=True` for everything), which is the first thing anyone opens when a UI test fails in CI.",
            "why": "When firmware 7.5 renames the Save button, you fix one line in the Page Object instead of forty tests. Screenshots turn \"it failed in CI\" into something you can see.",
            "syntax": "class AdminPage:\n    def __init__(self, page):\n        self.save = page.get_by_role(\"button\", name=\"Save\")\n    def set_channel(self, ch): ...\n\npage.screenshot(path=\"admin.png\")",
            "code": """from pathlib import Path  ## To read the page

from playwright.sync_api import expect, sync_playwright  ## Playwright sync API

HTML = Path("ap_admin.html").read_text()  ## Same admin UI as before


class AdminPage:  ## Page Object for the AP admin UI
    def __init__(self, page):  ## Wraps a Playwright page
        self.page = page  ## Keep a reference for navigation
        self.channel = page.get_by_label("Channel")  ## Locators are lazy, so defining them early is safe
        self.band = page.get_by_label("Band")  ## Dropdown
        self.message = page.locator("#msg")  ## Result line

    def open_and_login(self, user="admin", password="lab123"):  ## High-level action
        self.page.goto("http://ap.lab/")  ## Navigate
        self.page.get_by_label("Username").fill(user)  ## Credentials
        self.page.get_by_label("Password").fill(password)  ## Credentials
        self.page.get_by_role("button", name="Sign in").click()  ## Submit
        expect(self.page.get_by_text("Clients:")).to_be_visible()  ## Logged in when the status shows

    def set_channel(self, band, channel):  ## What the test case actually says
        self.band.select_option(band)  ## Pick the band
        self.channel.fill(str(channel))  ## Type the channel
        self.page.get_by_role("button", name="Save").click()  ## Save
        expect(self.message).to_have_text(f"Saved: {band} channel {channel}")  ## Confirmed by the UI


def route(r):  ## Fake device: canned status, accept saves, serve the page
    if "/api/" in r.request.url:  ## Any API call
        r.fulfill(json={"clients": 12, "uptime": "3d", "ok": True})  ## One canned reply fits both endpoints
    else:  ## The page itself
        r.fulfill(body=HTML, content_type="text/html")  ## Serve HTML


with sync_playwright() as p:  ## Driver lifecycle
    browser = p.chromium.launch()  ## Headless Chromium
    page = browser.new_page(viewport={"width": 640, "height": 260})  ## Fixed size for stable screenshots
    page.route("http://ap.lab/**", route)  ## Intercept the device
    admin = AdminPage(page)  ## The test talks to the Page Object only
    admin.open_and_login()  ## Step 1
    admin.set_channel("5G", 149)  ## Step 2, with its own verification
    page.screenshot(path="admin.png")  ## Evidence for the report
    print("screenshot bytes > 10 kB:", Path("admin.png").stat().st_size > 10_000)  ## It was written
    browser.close()  ## Done""",
            "under": "The screenshot below is the real file this code wrote. Locators created in `__init__` don't touch the page until an action or assertion uses them, so a Page Object can be built before navigation. Putting the verification inside `set_channel` means every caller gets it for free.",
            "gotchas": [],
            "best": "One Page Object per page or reusable component, methods named after user intent, no assertions about unrelated things inside them, a fixed viewport for screenshots, and screenshots, traces and videos saved only on failure (pytest-playwright's `--screenshot only-on-failure --tracing retain-on-failure`).",
        },
    ],
    "exercises": [
        {
            "level": "Easy",
            "local_only": PW,
            "prompt": "Load an HTML string with `page.set_content` and print the page title and the text of the `h1`.",
            "solution": "`set_content` loads HTML directly; `page.title()` and a role locator for the heading.",
            "solution_code": """from playwright.sync_api import sync_playwright

with sync_playwright() as p:
    page = p.chromium.launch().new_page()
    page.set_content("<title>Controller</title><h1>Site: Building 4</h1>")
    print(page.title())
    print(page.get_by_role("heading", level=1).text_content())""",
        },
        {
            "level": "Medium",
            "files": PW_FILES,
            "local_only": PW,
            "prompt": "Make the fake device return HTTP 500 for `/api/status` and check that the UI shows its error message.",
            "solution": "`route.fulfill(status=500, ...)` simulates a broken backend; `expect(...).to_contain_text` waits for the error banner.",
            "solution_code": """from pathlib import Path
from playwright.sync_api import expect, sync_playwright

HTML = Path("ap_admin.html").read_text()


def route(r):
    if r.request.url.endswith("/api/status"):
        r.fulfill(status=500, body="boom")
    else:
        r.fulfill(body=HTML, content_type="text/html")


with sync_playwright() as p:
    page = p.chromium.launch().new_page()
    page.route("http://ap.lab/**", route)
    page.goto("http://ap.lab/")
    page.get_by_role("button", name="Sign in").click()
    expect(page.locator("#status")).to_contain_text("status unavailable (500)")
    print(page.locator("#status").text_content(), "| class:", page.locator("#status").get_attribute("class"))""",
        },
        {
            "level": "Hard",
            "local_only": PW,
            "prompt": "A responsive page shows a hamburger menu button only below 600 px wide. Check it at a phone viewport (390×844) and a desktop viewport (1280×800).",
            "solution": "Create a browser context per viewport; `is_visible()` gives an instant answer (fine here because the CSS applies immediately).",
            "solution_code": """from playwright.sync_api import sync_playwright

HTML = '''<style>#menu{display:none} @media (max-width:600px){#menu{display:block} nav{display:none}}</style>
<button id="menu">Menu</button><nav>Dashboard | Radios | Clients</nav>'''

with sync_playwright() as p:
    browser = p.chromium.launch()
    for name, size in {"phone": (390, 844), "desktop": (1280, 800)}.items():
        ctx = browser.new_context(viewport={"width": size[0], "height": size[1]})
        page = ctx.new_page()
        page.set_content(HTML)
        print(f"{name:<8} menu visible: {page.get_by_role('button', name='Menu').is_visible()}")
        ctx.close()
    browser.close()""",
        },
    ],
    "interview": [
        {"level": "Intermediate", "q": "Playwright vs Selenium?", "a": "Playwright: auto-waiting actions and assertions, network interception, browser contexts for isolation, tracing, one API for Chromium/Firefox/WebKit, fast parallel runs. Selenium: W3C WebDriver standard, broad language and grid ecosystem, real-browser vendor drivers. For new projects Playwright usually gives less flaky tests.", "follow": "What's a browser context? (An isolated incognito-like session within one browser process.)"},
        {"level": "Intermediate", "q": "Why avoid time.sleep in UI tests?", "a": "Fixed sleeps are either too short (flaky) or too long (slow). Playwright's actions and expect assertions already wait for the right condition up to a timeout.", "follow": "How would you wait for a network response? (page.expect_response(url_pattern).)"},
        {"level": "Intermediate", "q": "What's the Page Object Model?", "a": "A design pattern where each page/component is a class exposing user-level actions and locators. Tests use those methods, so UI changes are fixed in one place and tests read like test cases.", "follow": "Should page objects contain assertions? (Only ones that confirm the action itself completed.)"},
        {"level": "Advanced", "q": "How would you test a device UI when the device is slow or unavailable in CI?", "a": "Serve the UI (or a recorded build) and intercept API calls with page.route, returning recorded JSON fixtures, including error and slow responses. Keep a smaller suite against real devices on the bench nightly.", "follow": "How do you record realistic fixtures? (HAR recording: context.route_from_har.)"},
    ],
    "quiz": [
        {"q": "Most robust locator for a button:", "options": ["#root > div:nth-child(3) button", "get_by_role('button', name='Save')", "XPath by index", "page.locator('button')"], "answer": 1, "why": "User-facing and stable."},
        {"q": "`expect(loc).to_have_text('x')` does what if text isn't there yet?", "options": ["Fails immediately", "Retries until timeout", "Sleeps 1 s once", "Skips"], "answer": 1, "why": "Auto-retrying assertion."},
        {"q": "Fake a device API in a UI test:", "options": ["page.route + route.fulfill", "monkeypatch", "Edit hosts file", "Not possible"], "answer": 0, "why": "Network interception."},
        {"q": "Default expect timeout:", "options": ["1 s", "5 s", "30 s", "None"], "answer": 1, "why": "5 seconds, configurable."},
        {"q": "Isolated session sharing one browser process:", "options": ["Page", "BrowserContext", "Frame", "Worker"], "answer": 1, "why": "Contexts isolate cookies and storage."},
    ],
    "cards": [
        {"front": "Install Playwright", "back": "`pip install playwright && playwright install chromium`"},
        {"front": "Locators to prefer", "back": "get_by_role, get_by_label, get_by_text, get_by_test_id"},
        {"front": "Assert with waiting", "back": "`expect(locator).to_have_text(...) / to_be_visible()`"},
        {"front": "Mock a backend", "back": "`page.route('**/api/**', lambda r: r.fulfill(json=...))`"},
        {"front": "Evidence on failure", "back": "`page.screenshot(path=...)`, tracing, video"},
    ],
}


M_6_4 = {
    "id": "6.4",
    "title": "Excel & CSV Reports",
    "hours": 2,
    "summary": "Read and write CSV with the standard library, summarise measurements, then build a formatted multi-sheet Excel report with pandas and openpyxl.",
    "concepts": [
        {
            "title": "csv.DictReader / DictWriter and summarising",
            "concept": "The `csv` module reads rows as lists (`csv.reader`) or as dicts keyed by the header (`csv.DictReader`), and writes them back with `csv.writer` / `csv.DictWriter`. Every value read from CSV is a **string**: convert numbers yourself. When writing real files, open them with `newline=\"\"` so the module controls line endings.",
            "why": "iperf wrappers, attenuator sweeps and test runners often emit CSV, and managers often want CSV back. The standard library handles quoting, commas inside values and headers correctly; `split(\",\")` doesn't.",
            "syntax": "with open(path, newline=\"\") as f:\n    for row in csv.DictReader(f): ...\n\nwith open(out, \"w\", newline=\"\") as f:\n    w = csv.DictWriter(f, fieldnames=[...])\n    w.writeheader(); w.writerow({...})",
            "code": """import csv  ## Standard-library CSV reader/writer
import io  ## In-memory text "files" for the example
import statistics  ## mean()

RAW = \"\"\"ap,band,channel,mbps
ap-1,5G,36,912.4
ap-1,5G,36,898.1
ap-2,6G,37,2410.0
ap-2,6G,37,2388.5
ap-3,2G,6,141.7
\"\"\"  ## What an iperf wrapper might have written

rows = list(csv.DictReader(io.StringIO(RAW)))  ## Each row becomes a dict keyed by the header
print(rows[0])  ## Note: every value is a string

groups = {}  ## (ap, band) -> list of measurements
for row in rows:  ## Group the rows
    groups.setdefault((row["ap"], row["band"]), []).append(float(row["mbps"]))  ## Convert explicitly

out = io.StringIO()  ## For a real file: open("summary.csv", "w", newline="")
fields = ["ap", "band", "runs", "mean_mbps", "min_mbps"]  ## Fixed column order
writer = csv.DictWriter(out, fieldnames=fields, lineterminator="\\n")  ## \\n endings for display; the default is \\r\\n
writer.writeheader()  ## Header line
for (ap, band), values in sorted(groups.items()):  ## Stable output order
    writer.writerow({"ap": ap, "band": band, "runs": len(values), "mean_mbps": round(statistics.mean(values), 1), "min_mbps": min(values)})  ## One summary row
print(out.getvalue())  ## The resulting CSV text""",
            "under": "The csv writer quotes fields that contain the delimiter, quotes or newlines, which is why it's safe for SSIDs like `Lab, 2nd floor`. The CSV standard (RFC 4180) uses `\\r\\n` line endings, the writer's default. Opening files without `newline=\"\"` lets Python translate line endings too, which produces blank lines between rows on Windows.",
            "gotchas": [
                {
                    "title": "Sorting or comparing CSV numbers as strings",
                    "note": "\"912.4\" > \"2410.0\" as strings, because comparison is character by character and '9' > '2'. Convert to numbers before sorting, comparing or summing.",
                    "wrong": """import csv
import io

rows = list(csv.DictReader(io.StringIO("ap,mbps\\nap-1,912.4\\nap-2,2410.0\\nap-3,141.7\\n")))
best = max(rows, key=lambda r: r["mbps"])
print("best:", best["ap"], best["mbps"])""",
                    "fixed": """import csv
import io

rows = list(csv.DictReader(io.StringIO("ap,mbps\\nap-1,912.4\\nap-2,2410.0\\nap-3,141.7\\n")))
best = max(rows, key=lambda r: float(r["mbps"]))
print("best:", best["ap"], best["mbps"])""",
                },
            ],
            "best": "Use DictReader/DictWriter (column order changes don't break you), convert types at the edge in one place, open files with `newline=\"\"` and an explicit `encoding=\"utf-8\"`, and switch to pandas once you need grouping, pivots or Excel.",
        },
        {
            "title": "pandas + openpyxl: a formatted Excel report",
            "local_only": PANDAS,
            "concept": "pandas loads records into a DataFrame, where `groupby(...).agg(...)` summarises and boolean columns flag KPI failures. `DataFrame.to_excel` with an `ExcelWriter` writes several sheets into one workbook. openpyxl then reopens it for formatting pandas doesn't do: bold headers, frozen panes, column widths and red fills for failing rows.",
            "why": "Results often go to people who live in Excel. A workbook with a Summary sheet (failures highlighted) and a Raw sheet answers both \"did we pass?\" and \"show me the numbers\".",
            "syntax": "summary = df.groupby([\"ap\", \"band\"], as_index=False).agg(mean=(\"mbps\", \"mean\"))\nwith pd.ExcelWriter(path, engine=\"openpyxl\") as xl:\n    summary.to_excel(xl, sheet_name=\"Summary\", index=False)",
            "code": """import pandas as pd  ## pip install pandas openpyxl
from openpyxl import load_workbook  ## To reopen and format the workbook
from openpyxl.styles import Font, PatternFill  ## Cell styles

runs = pd.DataFrame([  ## One record per measurement
    {"ap": "ap-1", "band": "5G", "mbps": 912.4}, {"ap": "ap-1", "band": "5G", "mbps": 840.0},  ## ap-1 dips below KPI once
    {"ap": "ap-2", "band": "6G", "mbps": 2410.0}, {"ap": "ap-2", "band": "6G", "mbps": 2388.5},  ## ap-2 is healthy
    {"ap": "ap-3", "band": "2G", "mbps": 141.7}, {"ap": "ap-3", "band": "2G", "mbps": 133.2},  ## ap-3 on 2.4 GHz
])  ## Closes the DataFrame call
KPI = {"2G": 120, "5G": 850, "6G": 2200}  ## Minimum acceptable throughput per band

summary = runs.groupby(["ap", "band"], as_index=False).agg(runs=("mbps", "size"), mean_mbps=("mbps", "mean"), min_mbps=("mbps", "min"))  ## Named aggregations
summary["mean_mbps"] = summary["mean_mbps"].round(1)  ## Tidy decimals
summary["kpi"] = summary["band"].map(KPI)  ## Look up each row's floor
summary["pass"] = summary["min_mbps"] >= summary["kpi"]  ## Vectorised comparison: a bool column
print(summary.to_string(index=False))  ## Console view

with pd.ExcelWriter("report.xlsx", engine="openpyxl") as xl:  ## One workbook, several sheets
    summary.to_excel(xl, sheet_name="Summary", index=False)  ## Summary first
    runs.to_excel(xl, sheet_name="Raw", index=False)  ## Raw data for drill-down

wb = load_workbook("report.xlsx")  ## Reopen to format
ws = wb["Summary"]  ## The summary sheet
ws.freeze_panes = "A2"  ## Header stays visible while scrolling
red = PatternFill("solid", start_color="FFC7CE")  ## Light red fill
for cell in ws[1]:  ## Header row (rows are 1-based)
    cell.font = Font(bold=True)  ## Bold headers
    ws.column_dimensions[cell.column_letter].width = 12  ## Readable column width
for row in ws.iter_rows(min_row=2):  ## Data rows
    if row[-1].value is False:  ## Last column is "pass"
        for cell in row:  ## Highlight the whole row
            cell.fill = red  ## Failing KPI
wb.save("report.xlsx")  ## Write the formatting back
print(wb.sheetnames, "failing rows:", sum(r[-1].value is False for r in ws.iter_rows(min_row=2)))  ## Verify""",
            "under": "`groupby` splits rows by key, applies the aggregations to each group and combines the results; named aggregation (`new=(column, func)`) controls the output column names. pandas writes Excel through openpyxl, which stores booleans as real TRUE/FALSE cells, so the check `value is False` works after reloading. openpyxl rows and columns are 1-based, like Excel.",
            "gotchas": [],
            "best": "Keep raw data alongside the summary, put thresholds in data (a KPI table) rather than in formatting code, use conditional formatting or fills for pass/fail, and write the file once at the end of the run so a crash doesn't leave half a report.",
        },
    ],
    "exercises": [
        {
            "level": "Easy",
            "prompt": "Write three channel-scan rows (ssid, channel, rssi) to `scan.csv` with `csv.writer`, then read them back and print the strongest SSID.",
            "solution": "Write with newline=\"\", read with DictReader, convert rssi to int when comparing.",
            "solution_code": """import csv

with open("scan.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["ssid", "channel", "rssi"])
    w.writerows([["lab-5g", 36, -52], ["guest", 6, -71], ["Lab, 2nd floor", 149, -48]])

with open("scan.csv", newline="", encoding="utf-8") as f:
    rows = list(csv.DictReader(f))
print(max(rows, key=lambda r: int(r["rssi"]))["ssid"])
print(open("scan.csv").read().splitlines()[3])""",
        },
        {
            "level": "Medium",
            "local_only": PANDAS,
            "prompt": "Build a pivot table of mean throughput with APs as rows and bands as columns from a list of measurements.",
            "solution": "`pivot_table(index=..., columns=..., values=..., aggfunc=\"mean\")`; missing combinations become NaN.",
            "solution_code": """import pandas as pd

df = pd.DataFrame({
    "ap": ["ap-1", "ap-1", "ap-2", "ap-2", "ap-2"],
    "band": ["5G", "6G", "5G", "6G", "6G"],
    "mbps": [910.0, 2300.0, 880.0, 2400.0, 2380.0],
})
print(df.pivot_table(index="ap", columns="band", values="mbps", aggfunc="mean"))""",
        },
        {
            "level": "Hard",
            "local_only": PANDAS,
            "prompt": "With openpyxl alone, create a workbook with per-AP throughput, an Excel formula row computing the average, and a bar chart. Reload it and confirm the formula text and the chart count.",
            "solution": "Formulas are stored as strings starting with `=` (Excel computes them when the file is opened); charts come from `openpyxl.chart`. After reloading, the cell still holds the formula text, not a number: openpyxl never calculates formulas (`load_workbook(data_only=True)` returns the value Excel cached the last time it saved the file, which is None for a file Excel never opened).",
            "solution_code": """from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference

wb = Workbook()
ws = wb.active
ws.title = "Throughput"
ws.append(["AP", "Mbps"])
for row in [("ap-1", 912), ("ap-2", 2410), ("ap-3", 142)]:
    ws.append(row)
ws["A5"], ws["B5"] = "Average", "=AVERAGE(B2:B4)"
chart = BarChart()
chart.title = "Throughput per AP"
chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
ws.add_chart(chart, "D2")
print("charts before save:", len(ws._charts))
wb.save("tput.xlsx")

ws2 = load_workbook("tput.xlsx")["Throughput"]
print("formula:", ws2["B5"].value)
print("charts after reload:", len(ws2._charts))""",
        },
    ],
    "interview": [
        {"level": "Beginner", "q": "Why use the csv module instead of split(',')?", "a": "It handles quoted fields containing commas, quotes and newlines, different dialects, and writes correctly escaped output. split breaks on `\"Lab, 2nd floor\"`.", "follow": "Why open CSV files with newline=''? (So the csv module controls line endings; avoids blank lines on Windows.)"},
        {"level": "Intermediate", "q": "How do you summarise measurements per AP and band in pandas?", "a": "df.groupby(['ap', 'band']).agg(mean=('mbps','mean'), min=('mbps','min'), runs=('mbps','size')), optionally as_index=False for a flat table.", "follow": "groupby vs pivot_table? (pivot_table reshapes one key into columns.)"},
        {"level": "Intermediate", "q": "pandas to_excel vs openpyxl directly?", "a": "to_excel is quickest for writing data frames; openpyxl gives full control of styles, formulas, charts, freeze panes and reading existing templates. Common pattern: write with pandas, format with openpyxl.", "follow": "Which library reads .xls (old format)? (xlrd; openpyxl only does .xlsx.)"},
    ],
    "quiz": [
        {"q": "Values from csv.DictReader are...", "options": ["Typed automatically", "Always strings", "Floats", "Bytes"], "answer": 1, "why": "Convert yourself."},
        {"q": "Open a CSV file for writing with...", "options": ["mode 'wb'", "newline=''", "encoding='ascii'", "buffering=0"], "answer": 1, "why": "Avoids doubled line endings."},
        {"q": "`max(['912.4', '2410.0'])` is...", "options": ["'2410.0'", "'912.4'", "Error", "2410.0"], "answer": 1, "why": "String comparison: '9' > '2'."},
        {"q": "Several sheets in one workbook with pandas:", "options": ["to_excel twice to the same path", "pd.ExcelWriter context", "to_csv", "Not possible"], "answer": 1, "why": "ExcelWriter."},
        {"q": "openpyxl row/column indexes are...", "options": ["0-based", "1-based", "Negative", "Letters only"], "answer": 1, "why": "Like Excel."},
    ],
    "cards": [
        {"front": "Read CSV as dicts", "back": "`csv.DictReader(open(p, newline='', encoding='utf-8'))`"},
        {"front": "Group and aggregate", "back": "`df.groupby(keys, as_index=False).agg(new=('col', 'mean'))`"},
        {"front": "Multi-sheet Excel", "back": "`with pd.ExcelWriter(p, engine='openpyxl') as xl: df.to_excel(xl, sheet_name=...)`"},
        {"front": "Format cells", "back": "openpyxl: Font(bold=True), PatternFill('solid', start_color=...), freeze_panes"},
        {"front": "Excel formula", "back": "`ws['B5'] = '=AVERAGE(B2:B4)'`"},
    ],
}


RVR_DATA = """ATTEN = list(range(0, 61, 6))
TPUT_5G = [935, 930, 921, 902, 870, 811, 720, 590, 410, 210, 40]
TPUT_6G = [2410, 2398, 2370, 2300, 2170, 1950, 1600, 1150, 640, 220, 0]
"""

M_6_5 = {
    "id": "6.5",
    "title": "Charts with matplotlib and plotly",
    "hours": 2,
    "summary": "Turn measurements into charts: a rate-vs-range plot with matplotlib (headless, saved to PNG for reports) and an interactive plotly chart saved as HTML.",
    "concepts": [
        {
            "title": "A rate-vs-range chart with matplotlib",
            "files": {"rvr_data.py": RVR_DATA},
            "local_only": MPL,
            "capture": "rvr.png",
            "image_alt": "Line chart of throughput versus attenuation for 5 GHz and 6 GHz, with a dashed KPI line at 500 Mbps",
            "concept": "matplotlib's object-oriented API: `fig, ax = plt.subplots()` creates a figure and axes, `ax.plot(x, y, marker=..., label=...)` draws a series, `ax.set_xlabel / set_ylabel / set_title / grid / legend` finish it, and `fig.savefig(path, dpi=..., bbox_inches=\"tight\")` writes a PNG. On servers select the non-GUI `Agg` backend and close each figure when done.",
            "why": "A rate-vs-range (RvR) chart is the single most-read output of a Wi-Fi performance run: it shows at a glance where each band falls off and whether it clears the KPI.",
            "syntax": "fig, ax = plt.subplots(figsize=(7, 4))\nax.plot(x, y, marker=\"o\", label=\"5 GHz\")\nax.axhline(kpi, linestyle=\"--\")\nfig.savefig(\"rvr.png\", dpi=100, bbox_inches=\"tight\")\nplt.close(fig)",
            "code": """import matplotlib  ## pip install matplotlib

matplotlib.use("Agg")  ## Non-GUI backend: required on CI machines without a display
import matplotlib.pyplot as plt  ## The plotting API (import after choosing the backend)

from rvr_data import ATTEN, TPUT_5G, TPUT_6G  ## Measured data (see supporting files)

KPI = 500  ## Minimum acceptable throughput, Mbps

fig, ax = plt.subplots(figsize=(7, 4))  ## One figure with one set of axes, size in inches
ax.plot(ATTEN, TPUT_5G, marker="o", label="5 GHz, 80 MHz")  ## First series
ax.plot(ATTEN, TPUT_6G, marker="s", label="6 GHz, 160 MHz")  ## Second series, different marker
ax.axhline(KPI, color="grey", linestyle="--", linewidth=1, label=f"KPI {KPI} Mbps")  ## Horizontal reference line
ax.set_xlabel("Attenuation (dB)")  ## Always label units
ax.set_ylabel("TCP throughput (Mbps)")  ## Y axis label
ax.set_title("Rate vs range, ap-lab-07, fw 7.4.1")  ## Title with the context people ask for
ax.grid(alpha=0.3)  ## Light grid for reading values
ax.legend()  ## Uses each series' label
fig.savefig("rvr.png", dpi=100, bbox_inches="tight")  ## PNG for the report; tight trims empty margins
plt.close(fig)  ## Free the figure's memory

for name, series in (("5G", TPUT_5G), ("6G", TPUT_6G)):  ## Same data, as numbers for the log
    last_ok = max(a for a, t in zip(ATTEN, series) if t >= KPI)  ## Highest attenuation still meeting KPI
    print(f"{name} meets {KPI} Mbps up to {last_ok} dB")  ## The headline numbers""",
            "under": "pyplot keeps a registry of open figures so `plt.show()` can display them; in scripts that only save files that registry is just memory you must release with `plt.close(fig)`. `Agg` renders with the Anti-Grain Geometry library into a raster buffer, and `savefig` encodes that buffer as PNG. The chart below is the actual file written by this code.",
            "gotchas": [
                {
                    "title": "Never closing figures in a loop",
                    "local_only": MPL,
                    "note": "Each `plt.subplots()` creates a figure that pyplot keeps alive until closed. A per-AP or per-test chart loop leaks memory, and matplotlib warns after 20 open figures. Close each one after saving.",
                    "wrong": """import warnings
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

with warnings.catch_warnings(record=True) as caught:
    warnings.simplefilter("always")
    for ap in range(25):
        fig, ax = plt.subplots()
        ax.plot([0, 1], [ap, ap])
print("open figures:", len(plt.get_fignums()))
print(caught[0].category.__name__, str(caught[0].message)[:40])""",
                    "fixed": """import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

for ap in range(25):
    fig, ax = plt.subplots()
    ax.plot([0, 1], [ap, ap])
    plt.close(fig)
print("open figures:", len(plt.get_fignums()))""",
                },
            ],
            "best": "Use the `fig, ax` API (not the implicit `plt.plot` state), label axes with units, put DUT, firmware and date in the title, use distinct markers as well as colours (colour-blind readers, greyscale prints), save at a fixed size and dpi, and close every figure.",
        },
        {
            "title": "Interactive charts with plotly",
            "files": {"rvr_data.py": RVR_DATA},
            "local_only": PLOTLY,
            "concept": "plotly builds interactive charts (hover values, zoom, toggling series) as a JSON figure spec rendered by JavaScript. `go.Figure` with `go.Scatter` traces gives full control; `plotly.express` is a one-line API for DataFrames. `fig.write_html(path, include_plotlyjs=\"cdn\")` produces a single HTML file anyone can open in a browser.",
            "why": "Dashboards and HTML reports benefit from hover-to-read values: an engineer can check the exact throughput at 42 dB without squinting at a PNG.",
            "syntax": "fig = go.Figure()\nfig.add_trace(go.Scatter(x=x, y=y, mode=\"lines+markers\", name=\"5 GHz\"))\nfig.update_layout(title=..., xaxis_title=..., yaxis_title=...)\nfig.write_html(\"rvr.html\", include_plotlyjs=\"cdn\")",
            "code": """from pathlib import Path  ## To inspect the output file

import plotly.graph_objects as go  ## pip install plotly

from rvr_data import ATTEN, TPUT_5G, TPUT_6G  ## Same measurements

fig = go.Figure()  ## Empty figure
fig.add_trace(go.Scatter(x=ATTEN, y=TPUT_5G, mode="lines+markers", name="5 GHz"))  ## One trace per series
fig.add_trace(go.Scatter(x=ATTEN, y=TPUT_6G, mode="lines+markers", name="6 GHz"))  ## Second trace
fig.add_hline(y=500, line_dash="dash", annotation_text="KPI 500 Mbps")  ## Reference line with a label
fig.update_layout(title="Rate vs range, ap-lab-07", xaxis_title="Attenuation (dB)", yaxis_title="Mbps", hovermode="x unified")  ## Axis titles and shared hover
fig.write_html("rvr.html", include_plotlyjs="cdn")  ## Small HTML file; plotly.js loads from a CDN

html = Path("rvr.html").read_text()  ## Check what was written
print("traces:", [t.name for t in fig.data])  ## Series in the figure
print("size under 20 kB:", len(html) < 20_000, "| uses CDN:", "cdn.plot.ly" in html)  ## Tiny because the JS isn't embedded
print(fig.to_dict()["layout"]["xaxis"]["title"]["text"])  ## The figure is plain data underneath""",
            "under": "A plotly figure is a dict of `data` (traces) and `layout`; Python only builds that JSON, and plotly.js draws it in the browser. `include_plotlyjs=\"cdn\"` references the library instead of embedding about 3.5 MB of JavaScript; use `True` for files that must work offline. Static images (PNG/SVG) need the extra `kaleido` package.",
            "gotchas": [],
            "best": "Use matplotlib PNGs for PDFs, emails and archived reports, plotly HTML for interactive dashboards. Keep the plotting code separate from the measurement code (functions that take data and return a figure) so it can be tested and reused.",
        },
    ],
    "exercises": [
        {
            "level": "Easy",
            "local_only": MPL,
            "capture": "bars.png",
            "image_alt": "Bar chart of mean throughput for three access points",
            "prompt": "Draw a bar chart of mean throughput per AP (ap-1 912, ap-2 2410, ap-3 142 Mbps) with value labels on the bars, and save it as bars.png.",
            "solution": "`ax.bar` returns the bar container, and `ax.bar_label` writes each value above its bar.",
            "solution_code": """import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

aps, mbps = ["ap-1", "ap-2", "ap-3"], [912, 2410, 142]
fig, ax = plt.subplots(figsize=(5, 3))
bars = ax.bar(aps, mbps, color=["#0A6C8A", "#3E9BB6", "#9CC9D8"])
ax.bar_label(bars, fmt="%d Mbps")
ax.set_ylabel("Mean throughput (Mbps)")
ax.set_ylim(0, 2800)
fig.savefig("bars.png", dpi=90, bbox_inches="tight")
plt.close(fig)
print("saved bars.png")""",
        },
        {
            "level": "Medium",
            "local_only": MPL,
            "capture": "roam.png",
            "image_alt": "Two stacked charts sharing a time axis: RSSI on top, throughput below, with a roam event marked",
            "prompt": "Plot RSSI and throughput over time in two stacked subplots that share the x axis, and mark a roam at t=12 s on both.",
            "solution": "`plt.subplots(2, 1, sharex=True)` returns two axes; `axvline` marks the event on each.",
            "solution_code": """import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

t = list(range(25))
rssi = [-50 - i * 2 if i < 12 else -48 - (i - 12) for i in t]
tput = [900 - i * 30 if i < 12 else 880 for i in t]
fig, (top, bottom) = plt.subplots(2, 1, sharex=True, figsize=(6, 4))
top.plot(t, rssi, color="#0A6C8A")
top.set_ylabel("RSSI (dBm)")
bottom.plot(t, tput, color="#C2410C")
bottom.set_ylabel("Mbps")
bottom.set_xlabel("Time (s)")
for ax in (top, bottom):
    ax.axvline(12, color="grey", linestyle=":")
top.annotate("roam to ap-2", (12, -60), xytext=(14, -66), arrowprops={"arrowstyle": "->"})
fig.savefig("roam.png", dpi=90, bbox_inches="tight")
plt.close(fig)
print("min RSSI before roam:", min(rssi[:12]))""",
        },
        {
            "level": "Hard",
            "local_only": MPL,
            "capture": "util.png",
            "image_alt": "Heatmap of channel utilisation per hour for six 5 GHz channels with a colour bar",
            "prompt": "Draw a heatmap of channel utilisation (%) for six 5 GHz channels over 8 hours with `imshow`, a colour bar, channel labels and each cell's value written in it.",
            "solution": "`imshow` maps a 2-D grid to colours; set tick labels from the channel list and annotate each cell with `ax.text`.",
            "solution_code": """import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

channels = [36, 52, 100, 116, 149, 165]
grid = [[(c * 7 + h * 11) % 90 + 5 for h in range(8)] for c in range(6)]
fig, ax = plt.subplots(figsize=(7, 3.5))
im = ax.imshow(grid, cmap="viridis", vmin=0, vmax=100, aspect="auto")
ax.set_yticks(range(6), [f"ch {c}" for c in channels])
ax.set_xticks(range(8), [f"{9 + h}:00" for h in range(8)])
for r, row in enumerate(grid):
    for c, v in enumerate(row):
        ax.text(c, r, v, ha="center", va="center", color="white" if v < 60 else "black", fontsize=8)
fig.colorbar(im, ax=ax, label="Utilisation (%)")
fig.savefig("util.png", dpi=90, bbox_inches="tight")
plt.close(fig)
print("busiest:", max((v, channels[r]) for r, row in enumerate(grid) for v in row))""",
        },
    ],
    "interview": [
        {"level": "Beginner", "q": "How do you make matplotlib work on a CI server with no display?", "a": "Use the Agg backend (matplotlib.use('Agg') before importing pyplot, or MPLBACKEND=Agg), save with savefig instead of show, and close figures.", "follow": "What happens if you call plt.show() with Agg? (Nothing is displayed; it's a no-op with a warning.)"},
        {"level": "Intermediate", "q": "pyplot state machine vs object-oriented API?", "a": "plt.plot() acts on an implicit current figure/axes, which gets confusing with several figures and in functions. fig, ax = plt.subplots() gives explicit objects to pass around, test and close.", "follow": "How do you get the current axes if you must? (plt.gca().)"},
        {"level": "Intermediate", "q": "When would you choose plotly over matplotlib?", "a": "For interactive output in browsers (hover, zoom, toggles) and dashboards. matplotlib is better for static, publication-quality images in PDFs and emails and has no JS dependency.", "follow": "How do you embed a plotly chart in your own HTML report? (fig.to_html(full_html=False, include_plotlyjs='cdn').)"},
    ],
    "quiz": [
        {"q": "Backend for headless servers:", "options": ["TkAgg", "Qt5Agg", "Agg", "MacOSX"], "answer": 2, "why": "Agg renders to files without a display."},
        {"q": "Free a figure's memory:", "options": ["del fig", "plt.close(fig)", "fig.clear()", "gc.collect()"], "answer": 1, "why": "pyplot keeps a reference until closed."},
        {"q": "Trim empty margins when saving:", "options": ["dpi=300", "bbox_inches='tight'", "transparent=True", "pad=0"], "answer": 1, "why": "Tight bounding box."},
        {"q": "Small plotly HTML file:", "options": ["include_plotlyjs=True", "include_plotlyjs='cdn'", "auto_open=True", "full_html=False"], "answer": 1, "why": "References the JS instead of embedding it."},
        {"q": "Two charts sharing an x axis:", "options": ["plt.twinx()", "plt.subplots(2, 1, sharex=True)", "ax.secondary", "Two figures"], "answer": 1, "why": "sharex links the axes."},
    ],
    "cards": [
        {"front": "Headless matplotlib", "back": "`matplotlib.use('Agg')` before `import matplotlib.pyplot`"},
        {"front": "Figure + axes", "back": "`fig, ax = plt.subplots(figsize=(7, 4))`"},
        {"front": "Save a chart", "back": "`fig.savefig('x.png', dpi=100, bbox_inches='tight'); plt.close(fig)`"},
        {"front": "Reference line", "back": "`ax.axhline(kpi, linestyle='--')` / `ax.axvline(x)`"},
        {"front": "Interactive HTML chart", "back": "`plotly` `go.Figure(...).write_html(p, include_plotlyjs='cdn')`"},
    ],
}
